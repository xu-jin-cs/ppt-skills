#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_check.py — ppt-direct 机考分项机械检查器（2026-08-15 裁定，gate-switch 实证族配套）

把四维×25 分门禁中可机械判定的分项（矢量合规 / 交付完整性）从模型自打分中剥离：
  - 文件存在
    （2026-08-17 用户裁定：不对 file_size 设定门槛——纯矢量红线与体积下限互斥，
      体积观测由 03_delivery.md 软层如实登记，不进机考）
  - slide 尺寸 13.333×7.5in
  - --expect-slides 给定时页数一致
  - 全部 shape 坐标红线：x+w ≤ 13.343in、y+h ≤ 7.51in、最小边 ≥ 0.5in
  - 文本 run 字号红线：全局最小 14pt
    （2026-08-17 用户裁定：删除 title 36~54 / body 14~24 上限——与 design_trace Hero≥88pt 闸
      双闸不可同真，择删上限保 Hero；title/body 粗分逻辑同步退役，美观上限由软层目检覆盖）
  - 图片类图表（PICTURE shape）数 == 0
  - pptx XML 中存在 a:ea typeface（中文字体红线）
  - --desktop-copy 给定时桌面副本存在

输出 JSON {pass, checks[{name, ok, detail}], violations}；任一违例 → exit 1，全过 → exit 0。
依赖：python-pptx（已装）。其余纯 stdlib。
"""
import argparse
import json
import os
import re
import sys
import zipfile

EMU_PER_IN = 914400
SLIDE_W_IN, SLIDE_H_IN = 13.333, 7.5
BOUND_W_IN, BOUND_H_IN = 13.343, 7.51
MIN_SIDE_IN = 0.5
# SIZE_MIN_B/SIZE_MAX_B 已随 2026-08-17 用户裁定删除（不对 file_size 设定门槛）
# TITLE_MIN/TITLE_MAX/BODY_MIN/BODY_MAX 已随 2026-08-17 用户裁定删除（Hero 双闸矛盾裁决：删上限保 Hero）
GLOBAL_MIN = 14

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
except ImportError:
    print(json.dumps({"pass": False, "violations": ["python-pptx 未安装"]},
                     ensure_ascii=False))
    sys.exit(1)


def _in(emu):
    return emu / EMU_PER_IN


def _is_title_shape(shape):
    """粗分 title/body：占位符 TITLE / 名称含 title|标题 / 顶部 1.0in 以内。"""
    try:
        if shape.is_placeholder and shape.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
            return True
    except Exception:
        pass
    name = (shape.name or "").lower()
    if "title" in name or "标题" in name:
        return True
    try:
        if shape.top is not None and _in(shape.top) < 1.0:
            return True
    except Exception:
        pass
    return False


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)


def main():
    ap = argparse.ArgumentParser(description="ppt-direct 机考分项机械检查器")
    ap.add_argument("--pptx", required=True, help="output.pptx 路径")
    ap.add_argument("--expect-slides", type=int, default=None, help="期望页数")
    ap.add_argument("--desktop-copy", default=None, help="桌面副本路径")
    args = ap.parse_args()

    checks, violations = [], []

    def rec(name, ok, detail):
        checks.append({"name": name, "ok": bool(ok), "detail": detail})
        if not ok:
            violations.append(f"{name}: {detail}")

    # 1. 文件存在（2026-08-17 裁定：体积不设门槛，仅记录观测值）
    if not os.path.isfile(args.pptx):
        rec("file_exists", False, f"{args.pptx} 不存在")
        return finish(checks, violations)
    rec("file_exists", True, args.pptx)
    size = os.path.getsize(args.pptx)
    rec("file_size", True, f"{size}B（观测值，不设门槛）")

    # 2. pptx 可解析
    try:
        prs = Presentation(args.pptx)
    except Exception as e:
        rec("pptx_parse", False, f"解析失败: {e}")
        return finish(checks, violations)
    rec("pptx_parse", True, "python-pptx 解析成功")

    # 3. slide 尺寸 13.333×7.5in
    w_in, h_in = _in(prs.slide_width), _in(prs.slide_height)
    rec("slide_size", abs(w_in - SLIDE_W_IN) < 0.01 and abs(h_in - SLIDE_H_IN) < 0.01,
        f"{w_in:.3f}×{h_in:.3f}in（要求 {SLIDE_W_IN}×{SLIDE_H_IN}in）")

    # 4. 页数
    n_slides = len(prs.slides)
    if args.expect_slides is not None:
        rec("slide_count", n_slides == args.expect_slides,
            f"{n_slides} 页（期望 {args.expect_slides}）")
    else:
        rec("slide_count", True, f"{n_slides} 页（未给定期望值，仅记录）")

    # 5. shape 坐标红线 + 6. 字号红线 + 7. 位图图表计数
    oob, font_bad, pictures = [], [], 0
    for si, slide in enumerate(prs.slides, 1):
        for sh in _iter_shapes(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            try:
                x, y = _in(sh.left), _in(sh.top)
                w, h = _in(sh.width), _in(sh.height)
            except Exception:
                continue
            if x + w > BOUND_W_IN + 1e-6 or y + h > BOUND_H_IN + 1e-6 \
                    or x < -1e-6 or y < -1e-6:
                oob.append(f"p{si}/{sh.name}: x+w={x + w:.3f} y+h={y + h:.3f}")
            if min(w, h) < MIN_SIDE_IN - 1e-6:
                oob.append(f"p{si}/{sh.name}: 最小边 {min(w, h):.3f}in < {MIN_SIDE_IN}in")
            if not sh.has_text_frame:
                continue
            # 2026-08-17 用户裁定：删除 title/body 字号上限判定（Hero 双闸矛盾，删上限保 Hero），
            # 仅保留全局最小 14pt 下限；_is_title_shape 保留不再参与判定
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is None or not run.text.strip():
                        continue
                    pt = run.font.size.pt
                    if pt < GLOBAL_MIN:
                        font_bad.append(f"p{si}/{sh.name}: {pt}pt < 全局最小 {GLOBAL_MIN}pt")
    rec("shape_bounds", not oob,
        "全部在界" if not oob else f"{len(oob)} 处越界: {'; '.join(oob[:5])}")
    rec("font_size", not font_bad,
        "字号全部合规" if not font_bad else f"{len(font_bad)} 处违例: {'; '.join(font_bad[:5])}")
    rec("bitmap_chart_count", pictures == 0, f"图片类图表 {pictures} 个（红线 0）")

    # 8. a:ea typeface 存在（解 pptx XML）
    ea_hit = False
    ea_re = re.compile(r"<a:ea\b[^>]*\btypeface=")
    try:
        with zipfile.ZipFile(args.pptx) as zf:
            for n in zf.namelist():
                if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
                    if ea_re.search(zf.read(n).decode("utf-8", "ignore")):
                        ea_hit = True
                        break
    except Exception as e:
        rec("ea_typeface", False, f"XML 解包失败: {e}")
    else:
        rec("ea_typeface", ea_hit, "存在 a:ea typeface" if ea_hit else "全部 slide XML 均无 a:ea typeface")

    # 9. 桌面副本
    if args.desktop_copy:
        rec("desktop_copy", os.path.isfile(args.desktop_copy),
            f"{args.desktop_copy} {'存在' if os.path.isfile(args.desktop_copy) else '不存在'}")
    else:
        rec("desktop_copy", True, "未给定 --desktop-copy，仅记录")

    return finish(checks, violations)


def finish(checks, violations):
    ok = not violations
    json.dump({"pass": ok, "checks": checks, "violations": violations},
              sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    # 末行单行摘要：gate_switch script_exit 原语取末行作 B 档违例详情
    summary = {"pass": ok, "violations": violations} if not ok else {"pass": ok}
    sys.stdout.write(json.dumps(summary, ensure_ascii=False, separators=(",", ":")) + "\n")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
