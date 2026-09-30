#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
design_trace_check.py — ppt-direct 设计溯源一致性机考（2026-08-17 扳手改造，gate-switch 实证族配套）

把四维×25 分门禁中「设计溯源一致性」分项从软打分剥离为机械判定（retro-pm-084 实证：
复刻成品四不像的真凶在 02 生成环节，色值/字号无人核验即进 03 目检）：

  ① 色板 binding 差集：提取产物 pptx 全部 run 色值（python-pptx run 级 + slide XML
     srgbClr 并集），与「01_design_spec 声明色（八色 HEX 等全部 #RRGGBB 词条）∪
     00_extracted_scheme 色板（fixed_theme_color 及文件内全部声明 HEX，复刻分支提供）
     ∪ --allow-colors 中性色豁免集」做差集；差集非空 = 自创 HEX 违例。
  ② Hero 焦点校验：产物全部 run 字号（python-pptx + XML sz 属性并集）最大值
     ≥ --hero-min-pt（默认 88）。
  ③ 涨跌徽章固定色号存在性：design_spec 中「涨/up」「降/down」行声明的 HEX（或
     --badge-colors 显式给定）必须出现在产物色值集合中；未声明则仅记录不判定。

用法：
  python3 design_trace_check.py --pptx <output.pptx> --design-spec <01_design_spec.md>
      [--extracted-scheme <00_extracted_scheme.json>] [--hero-min-pt 88]
      [--allow-colors FFFFFF,000000] [--badge-colors 2ED47A,FF7847]

退出码：0 = 全过 / 2 = 违例清单（末行单行 JSON 摘要供 gate_switch script_exit 摘取）。
路径与阈值全部参数化；默认豁免纯白/纯黑中性色，可用 --allow-colors "" 关闭。
依赖：python-pptx（已装）；其余纯 stdlib。
"""
import argparse
import json
import os
import re
import sys
import zipfile

# HEX 识别：#RRGGBB 必收；裸 6 位 HEX 仅当含数字才收（排除 decade/facade 类英文词误命中）
HEX_HASH_RE = re.compile(r"#([0-9A-Fa-f]{6})\b")
HEX_BARE_RE = re.compile(r"\b([0-9A-Fa-f]{6})\b")
SRGB_RE = re.compile(r'<a:srgbClr\s+val="([0-9A-Fa-f]{6})"')
SZ_RE = re.compile(r"\bsz=\"(\d+)\"")
UP_LINE_RE = re.compile(r"涨|\bup\b", re.IGNORECASE)
DOWN_LINE_RE = re.compile(r"降|\bdown\b", re.IGNORECASE)

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    print(json.dumps({"pass": False, "violations": ["python-pptx 未安装"]},
                     ensure_ascii=False))
    sys.exit(2)


def norm(hex6):
    return hex6.lstrip("#").upper()


def extract_hex_tokens(text):
    """从任意文本提取声明色集合（# 前缀 + 含数字的裸 6 位 HEX）。"""
    out = {norm(h) for h in HEX_HASH_RE.findall(text)}
    for tok in HEX_BARE_RE.findall(text):
        if any(c.isdigit() for c in tok):
            out.add(norm(tok))
    return out


def pptx_colors_and_sizes(path):
    """产物 pptx 色值集合 + 字号集合（python-pptx run 级 与 slide XML 并集）。"""
    colors, sizes = set(), set()
    prs = Presentation(path)

    def iter_shapes(shapes):
        for sh in shapes:
            yield sh
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from iter_shapes(sh.shapes)

    for slide in prs.slides:
        for sh in iter_shapes(slide.shapes):
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        sizes.add(run.font.size.pt)
                    try:
                        if run.font.color and run.font.color.rgb is not None:
                            colors.add(str(run.font.color.rgb).upper())
                    except Exception:
                        pass
    # slide XML 层：solidFill/line 等 run 外色值 + sz 属性兜底（仅 slides，不取 master/theme）
    with zipfile.ZipFile(path) as zf:
        for n in zf.namelist():
            if n.startswith("ppt/slides/slide") and n.endswith(".xml"):
                xml = zf.read(n).decode("utf-8", "ignore")
                colors.update(h.upper() for h in SRGB_RE.findall(xml))
                sizes.update(int(s) / 100.0 for s in SZ_RE.findall(xml))
    return colors, sizes


def parse_spec_badges(spec_text):
    """从 design_spec 行文本解析涨/降徽章声明色号（同色行 HEX）。"""
    up, down = set(), set()
    for line in spec_text.splitlines():
        hexes = extract_hex_tokens(line)
        if not hexes:
            continue
        if UP_LINE_RE.search(line):
            up.update(hexes)
        if DOWN_LINE_RE.search(line):
            down.update(hexes)
    return up, down


def main():
    ap = argparse.ArgumentParser(description="ppt-direct 设计溯源一致性机考")
    ap.add_argument("--pptx", required=True, help="产物 output.pptx 路径")
    ap.add_argument("--design-spec", required=True, help="01_design_spec.md 路径")
    ap.add_argument("--extracted-scheme", default="",
                    help="00_extracted_scheme.json 路径（复刻分支提供；直出分支留空）")
    ap.add_argument("--hero-min-pt", type=float, default=88.0,
                    help="Hero 大数字最小字号（pt）")
    ap.add_argument("--allow-colors", default="FFFFFF,000000",
                    help="中性色豁免集（逗号分隔 HEX；传空串关闭豁免）")
    ap.add_argument("--badge-colors", default="",
                    help="显式声明涨跌徽章色号（逗号分隔 HEX；缺省时从 design_spec 涨/降行解析）")
    args = ap.parse_args()

    checks, violations = [], []

    def rec(name, ok, detail):
        checks.append({"name": name, "ok": bool(ok), "detail": detail})
        if not ok:
            violations.append(f"{name}: {detail}")

    # 1. 输入存在性
    if not os.path.isfile(args.pptx):
        rec("pptx_exists", False, f"{args.pptx} 不存在")
        return finish(checks, violations)
    rec("pptx_exists", True, args.pptx)
    if not os.path.isfile(args.design_spec):
        rec("design_spec_exists", False, f"{args.design_spec} 不存在")
        return finish(checks, violations)
    rec("design_spec_exists", True, args.design_spec)

    spec_text = open(args.design_spec, encoding="utf-8", errors="ignore").read()

    # 2. 允许色集合 = design_spec 声明色 ∪ scheme 色板 ∪ 中性色豁免
    allowed = extract_hex_tokens(spec_text)
    scheme_path = (args.extracted_scheme or "").strip()
    scheme_colors = set()
    if scheme_path:
        if not os.path.isfile(scheme_path):
            rec("scheme_exists", False, f"{scheme_path} 不存在（复刻分支必传）")
            return finish(checks, violations)
        scheme_colors = extract_hex_tokens(
            open(scheme_path, encoding="utf-8", errors="ignore").read())
        allowed |= scheme_colors
        rec("scheme_palette", True,
            f"scheme 色板 {len(scheme_colors)} 色并入允许集: {sorted(scheme_colors)}")
    else:
        rec("scheme_palette", True, "直出分支，无 scheme（仅 design_spec 声明色）")
    allow_extra = {norm(h) for h in args.allow_colors.split(",") if h.strip()}
    allowed |= allow_extra
    rec("allowed_palette", True,
        f"允许色 {len(allowed)} 色（含豁免 {sorted(allow_extra)}）: {sorted(allowed)}")

    # 3. 产物色值/字号提取
    try:
        used_colors, used_sizes = pptx_colors_and_sizes(args.pptx)
    except Exception as e:
        rec("pptx_parse", False, f"解析失败: {e}")
        return finish(checks, violations)
    rec("pptx_parse", True,
        f"产物色值 {len(used_colors)} 色 / 字号 {len(used_sizes)} 档")

    # 4. ① 色板 binding 差集（自创 HEX 检测）
    rogue = sorted(used_colors - allowed)
    rec("color_binding", not rogue,
        "全部色值可溯源" if not rogue
        else f"自创 HEX 违例 {len(rogue)} 色（不在 design_spec/scheme/豁免集）: {rogue}")

    # 5. ② Hero 焦点 ≥ hero-min-pt
    max_pt = max(used_sizes) if used_sizes else 0.0
    rec("hero_min_size", used_sizes and max_pt >= args.hero_min_pt,
        f"最大字号 {max_pt:g}pt（要求 ≥{args.hero_min_pt:g}pt Hero 大数字）")

    # 6. ③ 涨跌徽章固定色号存在性（声明才判定）
    if args.badge_colors.strip():
        badge_declared = {norm(h) for h in args.badge_colors.split(",") if h.strip()}
        badge_src = "--badge-colors 显式声明"
    else:
        up, down = parse_spec_badges(spec_text)
        badge_declared = up | down
        badge_src = f"design_spec 涨/降行解析（涨 {sorted(up)} 降 {sorted(down)}）"
    if badge_declared:
        missing = sorted(badge_declared - used_colors)
        rec("badge_colors", not missing,
            f"徽章色号全部落地（{badge_src}）" if not missing
            else f"徽章色号缺失 {missing}（{badge_src}，产物未出现）")
    else:
        rec("badge_colors", True, "design_spec 未声明徽章色号，仅记录不判定")

    return finish(checks, violations)


def finish(checks, violations):
    ok = not violations
    json.dump({"pass": ok, "checks": checks, "violations": violations},
              sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    # 末行单行摘要：gate_switch script_exit 原语取末行作 B 档违例详情
    summary = {"pass": ok, "violations": violations} if not ok else {"pass": ok}
    sys.stdout.write(json.dumps(summary, ensure_ascii=False, separators=(",", ":")) + "\n")
    sys.exit(0 if ok else 2)


if __name__ == "__main__":
    main()
