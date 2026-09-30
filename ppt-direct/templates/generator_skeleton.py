# -*- coding: utf-8 -*-
"""ppt-direct 生成骨架 V1.0 — python-pptx 原生矢量绘制 helper 函数族
用法：from generator_skeleton import *  （或内联复制本文件全部 helper）
配套技能：/ppt-direct（节点 direct_generator 必须复用本骨架，禁止裸写 add_shape）
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------- 默认设计令牌（design_planner 可在 01_design_spec.md 中覆盖） ----------
NAVY   = "141B3D"   # 深底
NAVY2  = "1D2A5C"   # 深底次层
BLUE   = "4C6FFF"   # 主色
BLUE_D = "3F51B5"   # 主色深
ORANGE = "F5A623"   # 点缀
GREEN  = "2ED47A"   # 涨
RED_O  = "FF7847"   # 降（优化）
WHITE  = "FFFFFF"
INK    = "23283B"   # 浅底正文
MUTE   = "8A92A6"   # 弱化
BG_L   = "F5F7FC"   # 浅底
GRAY_L = "C9CFDB"   # 基线条
FONT   = "Microsoft YaHei"

SLIDE_W, SLIDE_H = 13.333, 7.5


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    return prs, prs.slide_layouts[6]


def hx(c):
    return RGBColor.from_string(c)


def set_alpha(shape, alpha):
    """python-pptx 无原生透明度 API，走 XML 注入"""
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf.find(qn("a:srgbClr"))
    if clr is None:
        return
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))


def rect(slide, x, y, w, h, color, alpha=None, line=None, round_=False, radius=0.12):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if round_:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = hx(color)
        if alpha is not None:
            set_alpha(shp, alpha)
    if line:
        shp.line.color.rgb = hx(line); shp.line.width = Pt(1.2)
    else:
        shp.line.fill.background()
    return shp


def circle(slide, x, y, d, color=None, alpha=None, line=None, lw=1.5):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.shadow.inherit = False
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = hx(color)
        if alpha is not None:
            set_alpha(shp, alpha)
    if line:
        shp.line.color.rgb = hx(line); shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    return shp


def text(slide, x, y, w, h, s, size, color, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.0):
    """统一文本框：自动设置中文字体（a:ea typeface）"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, ln in enumerate(str(s).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        r = p.add_run(); r.text = ln
        f = r.font; f.name = FONT; f.size = Pt(size); f.bold = bold; f.color.rgb = hx(color)
        rPr = r._r.get_or_add_rPr()
        rPr.append(rPr.makeelement(qn("a:ea"), {"typeface": FONT}))
    return tb


def pill(slide, x, y, w, h, s, bg, fg, size=12, bold=True):
    """胶囊标签/徽章"""
    rect(slide, x, y, w, h, bg, round_=True, radius=0.5)
    text(slide, x, y - 0.02, w, h, s, size, fg, bold=bold, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def page_bg(slide, color):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, color)


def header(slide, part, title, sub=""):
    """浅底数据页页眉：PART 标签 + 大标题 + accent 短线 + 副注"""
    text(slide, 0.8, 0.45, 3.0, 0.35, part, 12, BLUE, bold=True)
    text(slide, 0.8, 0.78, 11.0, 0.65, title, 27, INK, bold=True)
    rect(slide, 0.82, 1.5, 0.62, 0.055, ORANGE)
    if sub:
        text(slide, 0.8, 1.62, 11.7, 0.35, sub, 12, MUTE)


def footer(slide, dark=False, note="Generated with /ppt-direct", date=""):
    c = "5A6480" if dark else MUTE
    text(slide, 0.8, 7.08, 8.0, 0.35, note, 9, c)
    text(slide, 10.9, 7.08, 1.63, 0.35, date, 9, c, align=PP_ALIGN.RIGHT)


def divider(prs_slide, no, title, sub):
    """深底章节页：120pt 大序号 + PART 标签 + accent 短线 + 标题 + 副注"""
    s = prs_slide
    page_bg(s, NAVY)
    circle(s, 10.3, 3.9, 5.6, color=BLUE, alpha=0.12)
    circle(s, 11.5, 5.1, 3.2, color=None, line=BLUE, lw=1.2)
    text(s, 0.9, 1.5, 5.0, 2.2, no, 120, BLUE, bold=True)
    text(s, 0.95, 3.85, 4.0, 0.4, f"PART {no}", 14, ORANGE, bold=True)
    rect(s, 0.98, 4.32, 1.2, 0.06, ORANGE)
    text(s, 0.9, 4.5, 11.5, 0.95, title, 40, WHITE, bold=True)
    text(s, 0.95, 5.55, 11.0, 0.45, sub, 15, "AAB4D8")
    return s


def hbar_rows(slide, rows, top=2.15, row_h=1.02, max_idx=None,
              label_w=2.5, bar_x=3.55, bar_max_w=6.6):
    """K2 基线 vs 新值 双条矢量数据行 + 涨跌徽章
    rows: [(指标名, 基线标签, 新值标签, 新值指数(基线=100), 变化文案, 'up'|'down')]
    """
    max_idx = max_idx or max(r[3] for r in rows)
    scale = bar_max_w / max_idx
    for i, (name, v2, v3, idx, chg, direction) in enumerate(rows):
        y = top + i * row_h
        text(slide, 0.8, y - 0.06, label_w, 0.62, name, 13.5, INK, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        w2 = 100 * scale
        rect(slide, bar_x, y, max(w2, 0.12), 0.26, GRAY_L, round_=True, radius=0.5)
        text(slide, bar_x + max(w2, 0.12) + 0.08, y - 0.035, 1.6, 0.32, v2, 10.5, MUTE, anchor=MSO_ANCHOR.MIDDLE)
        w3 = idx * scale
        rect(slide, bar_x, y + 0.36, max(w3, 0.12), 0.26, BLUE, round_=True, radius=0.5)
        text(slide, bar_x + max(w3, 0.12) + 0.08, y + 0.325, 1.9, 0.32, v3, 11, BLUE_D, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        pill(slide, 11.15, y + 0.1, 1.38, 0.44, chg, GREEN if direction == "up" else RED_O, WHITE, size=13)
